import io
from datetime import date, datetime

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession

from app.database import get_db
from app.models import Receita

router = APIRouter(prefix="/api/relatorios", tags=["Relatorios"])

_CABECALHO = PatternFill("solid", fgColor="7C3AED")
_FONTE_CABECALHO = Font(color="FFFFFF", bold=True)


def _planilha(nome_aba: str, colunas: list[str], linhas: list[tuple]) -> io.BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = nome_aba[:31]

    ws.append(colunas)
    for celula in ws[1]:
        celula.fill = _CABECALHO
        celula.font = _FONTE_CABECALHO
        celula.alignment = Alignment(horizontal="center")

    for linha in linhas:
        ws.append(list(linha))

    for i, coluna in enumerate(colunas, start=1):
        largura = max(len(str(coluna)), *(len(str(l[i - 1])) for l in linhas)) if linhas else len(coluna)
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = min(largura + 4, 45)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def _excel_response(buffer: io.BytesIO, nome_arquivo: str) -> StreamingResponse:
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{nome_arquivo}"'},
    )


@router.get("/vendas.xlsx")
async def vendas_excel(
    inicio: date | None = None,
    fim: date | None = None,
    db: AsyncSession = Depends(get_db),
):
    sql = """
        SELECT venda_id, data_venda, cliente, produto, quantidade,
               preco_unitario, custo_unitario, subtotal, lucro_item,
               forma_pagamento, status
        FROM vw_vendas_detalhe
        WHERE 1 = 1
    """
    params: dict = {}
    if inicio:
        sql += " AND data_venda >= :inicio"
        params["inicio"] = inicio
    if fim:
        sql += " AND data_venda < (:fim + INTERVAL '1 day')"
        params["fim"] = fim
    sql += " ORDER BY data_venda DESC, venda_id"

    linhas = (await db.execute(text(sql), params)).all()

    dados = [
        (
            l.venda_id,
            l.data_venda.strftime("%d/%m/%Y %H:%M"),
            l.cliente or "-",
            l.produto or "-",
            float(l.quantidade),
            float(l.preco_unitario),
            float(l.custo_unitario),
            float(l.subtotal),
            float(l.lucro_item),
            l.forma_pagamento,
            l.status,
        )
        for l in linhas
    ]

    buffer = _planilha(
        "Vendas",
        ["Venda", "Data", "Cliente", "Produto", "Qtd", "Preco unit.",
         "Custo unit.", "Subtotal", "Lucro", "Pagamento", "Status"],
        dados,
    )
    return _excel_response(buffer, f"vendas_{datetime.now():%Y%m%d}.xlsx")


@router.get("/estoque.xlsx")
async def estoque_excel(db: AsyncSession = Depends(get_db)):
    linhas = (
        await db.execute(text("SELECT * FROM vw_estoque WHERE ativo ORDER BY categoria, nome"))
    ).all()

    dados = [
        (
            l.nome,
            l.categoria or "-",
            float(l.quantidade_atual),
            l.unidade_medida,
            float(l.estoque_minimo),
            float(l.custo_unitario),
            float(l.valor_em_estoque),
            l.situacao,
        )
        for l in linhas
    ]

    buffer = _planilha(
        "Estoque",
        ["Ingrediente", "Categoria", "Qtd atual", "Unidade", "Estoque min.",
         "Custo unit.", "Valor em estoque", "Situacao"],
        dados,
    )
    return _excel_response(buffer, f"estoque_{datetime.now():%Y%m%d}.xlsx")


@router.get("/lista-compras")
async def lista_compras(
    cobertura: float = Query(2.0, gt=0, description="Multiplo do estoque minimo a repor"),
    db: AsyncSession = Depends(get_db),
):
    """O que precisa comprar: insumos abaixo do minimo e quanto custa repor."""
    linhas = (
        await db.execute(
            text(
                """
                SELECT nome, categoria, unidade_medida, quantidade_atual,
                       estoque_minimo, custo_unitario, situacao, fornecedor
                FROM vw_estoque
                WHERE ativo AND situacao IN ('sem_estoque', 'baixo', 'atencao')
                ORDER BY situacao, nome
                """
            )
        )
    ).mappings().all()

    itens = []
    total = 0.0
    for l in linhas:
        alvo = float(l["estoque_minimo"]) * cobertura
        comprar = max(alvo - float(l["quantidade_atual"]), 0)
        custo = comprar * float(l["custo_unitario"])
        total += custo
        itens.append(
            {
                "ingrediente": l["nome"],
                "categoria": l["categoria"],
                "unidade": l["unidade_medida"],
                "quantidade_atual": float(l["quantidade_atual"]),
                "estoque_minimo": float(l["estoque_minimo"]),
                "comprar": round(comprar, 3),
                "custo_estimado": round(custo, 2),
                "situacao": l["situacao"],
                "fornecedor": l["fornecedor"],
            }
        )

    return {"itens": itens, "custo_total_estimado": round(total, 2), "cobertura": cobertura}


@router.get("/ficha-tecnica/{receita_id}.pdf")
async def ficha_tecnica_pdf(receita_id: int, db: AsyncSession = Depends(get_db)):
    receita = await db.get(Receita, receita_id)
    if not receita:
        raise HTTPException(404, "Receita nao encontrada")

    linhas = (
        await db.execute(
            text(
                """
                SELECT ingrediente, quantidade, unidade, custo_por_unidade,
                       custo_linha, observacao
                FROM vw_ficha_tecnica
                WHERE receita_id = :id
                ORDER BY custo_linha DESC
                """
            ),
            {"id": receita_id},
        )
    ).all()

    empresa = await db.scalar(
        text("SELECT valor FROM configuracoes WHERE chave = 'empresa_nome'")
    )

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=A4,
        leftMargin=2 * cm, rightMargin=2 * cm, topMargin=2 * cm, bottomMargin=2 * cm,
        title=f"Ficha tecnica - {receita.nome}",
    )
    estilos = getSampleStyleSheet()
    roxo = colors.HexColor("#7C3AED")

    elementos = [
        Paragraph(f"<b>{empresa or 'Confeitaria'}</b>", estilos["Title"]),
        Paragraph(f"Ficha tecnica: <b>{receita.nome}</b>", estilos["Heading2"]),
        Spacer(1, 6),
        Paragraph(
            f"Rendimento: {float(receita.rendimento):g} {receita.unidade_rendimento} &nbsp;|&nbsp; "
            f"Tempo de preparo: {receita.tempo_preparo_min or '-'} min &nbsp;|&nbsp; "
            f"Emitida em {datetime.now():%d/%m/%Y}",
            estilos["Normal"],
        ),
        Spacer(1, 14),
    ]

    tabela = [["Ingrediente", "Qtd", "Un.", "Custo/un.", "Custo"]]
    for l in linhas:
        tabela.append(
            [
                l.ingrediente,
                f"{float(l.quantidade):g}",
                l.unidade,
                f"R$ {float(l.custo_por_unidade):.6f}",
                f"R$ {float(l.custo_linha):.2f}",
            ]
        )

    t = Table(tabela, colWidths=[7.5 * cm, 2 * cm, 1.5 * cm, 3 * cm, 3 * cm])
    t.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), roxo),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D4D4D8")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F4F4F5")]),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
            ]
        )
    )
    elementos += [t, Spacer(1, 16)]

    custo_total = float(receita.custo_total)
    custo_unitario = float(receita.custo_unitario)
    preco = float(receita.preco_venda_sugerido or 0)
    resumo = [
        ["Custo total da receita", f"R$ {custo_total:.2f}"],
        [f"Custo por {receita.unidade_rendimento[:-1] if receita.unidade_rendimento.endswith('s') else receita.unidade_rendimento}",
         f"R$ {custo_unitario:.2f}"],
        ["Preco de venda sugerido", f"R$ {preco:.2f}" if preco else "-"],
        [
            "Margem de lucro",
            f"{(preco - custo_unitario) / preco * 100:.1f}%" if preco > 0 else "-",
        ],
    ]
    tr = Table(resumo, colWidths=[10 * cm, 7 * cm])
    tr.setStyle(
        TableStyle(
            [
                ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 10),
                ("ALIGN", (1, 0), (1, -1), "RIGHT"),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D4D4D8")),
                ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#EDE9FE")),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
            ]
        )
    )
    elementos.append(tr)

    if receita.descricao:
        elementos += [Spacer(1, 14), Paragraph(f"<b>Modo de preparo / observacoes</b><br/>{receita.descricao}", estilos["Normal"])]

    doc.build(elementos)
    buffer.seek(0)

    nome = receita.nome.lower().replace(" ", "_")
    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": f'inline; filename="ficha_{nome}.pdf"'},
    )
